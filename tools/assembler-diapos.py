#!/usr/bin/env python3
"""Assemble les diapos selon slides-src/ordre.txt :
images du site (slides/), miniatures, PDF et PowerPoint (presentation/)."""
import pathlib, shutil, sys
from PIL import Image

root = pathlib.Path(__file__).resolve().parent.parent
src = root / 'slides-src'
out = root / 'slides'
thumbs = out / 'thumbs'
pres = root / 'presentation'

files = []
for line in (src / 'ordre.txt').read_text(encoding='utf-8').splitlines():
    line = line.strip()
    if not line or line.startswith('#'):
        continue
    kind, name = line.split(None, 1)
    f = src / ('originales/page-%s.jpg' % name if kind == 'page' else 'rendu/%s.jpg' % name)
    if not f.exists():
        sys.exit('Introuvable : %s' % f)
    files.append(f)

thumbs.mkdir(parents=True, exist_ok=True)
pres.mkdir(exist_ok=True)
for old in list(out.glob('slide-*.jpg')) + list(thumbs.glob('slide-*.jpg')):
    old.unlink()
for i, f in enumerate(files, 1):
    dst = out / ('slide-%02d.jpg' % i)
    shutil.copyfile(f, dst)
    im = Image.open(f).convert('RGB')
    im.resize((400, 225), Image.LANCZOS).save(thumbs / dst.name, quality=78)

# PDF (secours / impression)
imgs = [Image.open(out / ('slide-%02d.jpg' % i)).convert('RGB') for i in range(1, len(files) + 1)]
imgs[0].save(pres / 'presentation.pdf', save_all=True, append_images=imgs[1:], resolution=144)

# PowerPoint (une image plein cadre par diapo, notes de l'orateur vides)
try:
    from pptx import Presentation
    from pptx.util import Emu
    p = Presentation()
    p.slide_width, p.slide_height = Emu(12192000), Emu(6858000)
    for i in range(1, len(files) + 1):
        s = p.slides.add_slide(p.slide_layouts[6])
        s.shapes.add_picture(str(out / ('slide-%02d.jpg' % i)), 0, 0, p.slide_width, p.slide_height)
    p.save(pres / 'presentation.pptx')
except ImportError:
    print('python-pptx absent : PowerPoint non généré (pip install python-pptx)')

print('%d diapos -> mettez slideCount: %d dans js/deck.js' % (len(files), len(files)))
